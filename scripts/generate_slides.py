import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Color definitions (Premium Navy & Steel Blue theme)
NAVY = RGBColor(30, 58, 138)
STEEL_BLUE = RGBColor(70, 130, 180)
CHARCOAL = RGBColor(40, 40, 40)
WHITE = RGBColor(255, 255, 255)

SLIDES_DATA = [
    {
        "title": "Mutual Fund Analytics Platform",
        "subtitle": "End-to-End Data Engineering & Interactive Dashboard\nPrepared by: Intern / Data Analyst — Bluestock Fintech\nJune 2026",
        "bullets": [],
        "is_title": True
    },
    {
        "title": "Problem Statement",
        "bullets": [
            "Data Fragmentation: NAV, AUM, and transactions are scattered across different platforms in varying formats.",
            "Performance Comparison Gap: Hard to compare multiple funds across different AMCs on a risk-adjusted basis.",
            "No Benchmark Tracking: Retail investors struggle to monitor if their funds are outperforming benchmark indices.",
            "Solution: A unified SQLite database with automated metrics calculations & dashboards."
        ]
    },
    {
        "title": "Project Objectives",
        "bullets": [
            "Build an automated ETL pipeline to ingest raw AMFI and API data.",
            "Design a normalized star schema database for mutual fund analytics.",
            "Compute key risk-adjusted performance metrics: Sharpe, Sortino, Alpha, Beta, Max Drawdown.",
            "Perform investor transaction cohort and retention continuation analysis.",
            "Deliver an interactive web dashboard for fund analysis."
        ]
    },
    {
        "title": "Data Ingestion & Inflow API",
        "bullets": [
            "Source 1: AMFI India published scheme master list (01_fund_master.csv).",
            "Source 2: mfapi.in REST API for daily historical NAVs of selected schemes.",
            "Source 3: NSE/BSE public indices for Nifty 50 and Nifty 100 historical prices.",
            "Source 4: Internal transaction data (32,000+ rows) and portfolio holdings.",
            "Validation: 100% of schemes in the master list were verified against the NAV history database."
        ]
    },
    {
        "title": "System Architecture",
        "bullets": [
            "Extract: Automated Python requests scripts pulling from mfapi.in daily.",
            "Transform: Pandas preprocessing (parsing dates, forward-filling weekend and holiday NAV gaps, calculating daily returns).",
            "Load: Normalized Star Schema loaded into SQLite (dim_fund, dim_date, fact_nav, fact_transactions, fact_performance).",
            "Analyze: Metrics engine computing statistical parameters.",
            "Visualize: Python Streamlit application hosting interactive dashboard views."
        ]
    },
    {
        "title": "Database Star Schema DDL",
        "bullets": [
            "dim_fund: AMFI code (PK), fund house, category, manager, expense ratio, risk grade.",
            "dim_date: Date key, year, month, quarter, weekday flag.",
            "fact_nav: Scheme code, date key, Net Asset Value (NAV), daily returns.",
            "fact_transactions: Tx ID (PK), investor ID, date key, transaction type, amount, state, city tier, age group, KYC status."
        ]
    },
    {
        "title": "Exploratory Data Analysis (EDA)",
        "bullets": [
            "AUM Concentration: SBI Mutual Fund maintains largest share (Rs. 12.5 lakh crore as of Dec 2025).",
            "SIP Milestone: Monthly SIP inflows peaked in December 2025 at Rs. 31,002 Crore.",
            "Investor Geographics: Top investment contributions driven by Maharashtra, Gujarat, Punjab, and Karnataka.",
            "Transactions: T30 cities represent 66% of total investment counts."
        ]
    },
    {
        "title": "Performance Scorecard Rankings",
        "bullets": [
            "Methodology: Composite Score (0-100) combining 30% 3yr returns, 25% Sharpe ratio, 20% Alpha, 15% lower expense ratio, and 10% lower Max Drawdown.",
            "Axis Midcap Fund: Top ranked scheme with a Sharpe ratio of 1.76 and annualized returns of 35.3%.",
            "HDFC Mid-Cap Opportunities Fund: Second ranked with a Sharpe ratio of 1.62."
        ]
    },
    {
        "title": "Advanced Risk Analytics",
        "bullets": [
            "Value at Risk (VaR 95%): Daily maximum loss ceiling calculated around 1.8% to 2.2% for standard equity schemes.",
            "Conditional VaR (CVaR): Identifies the average loss during worst 5% of trading days.",
            "Sector Concentration (HHI): Index portfolio analysis shows Banking, IT, and Pharmaceuticals as highly concentrated sectors."
        ]
    },
    {
        "title": "Investor Analytics & Cohorts",
        "bullets": [
            "Age Bracket: Investors aged 26-35 and 36-45 represent over 58% of total portfolios.",
            "SIP Gaps: 8% of active SIP investors identified with transaction gaps > 35 days and flagged as 'At-Risk'.",
            "KYC verification: 92% verified status overall, showing high compliance."
        ]
    },
    {
        "title": "Streamlit Interactive Dashboard",
        "bullets": [
            "Industry Overview Tab: Interactive KPI cards, AUM bar charts, and monthly SIP trends.",
            "Performance Scorecard Tab: Scatter plot return vs. risk and complete sortable ranks.",
            "Investor Analytics Tab: Geographic heatmaps, income vs. age distributions.",
            "Market Trends Tab: Category inflows and dual axis Nifty vs. SIP tracking."
        ]
    },
    {
        "title": "Conclusion & Scope",
        "bullets": [
            "Pipeline: Created a fully automated Python execution pipeline handling 100K+ combined rows.",
            "Database: Robust relational structure avoiding duplication with index optimizations.",
            "Visuals: Premium interactive dashboard allowing self-service analytical querying.",
            "Future Scope: Integrate Markowitz Efficient Frontier optimization module and automated weekly HTML email alerts."
        ]
    }
]

def create_deck():
    prs = Presentation()
    
    # Set slide dimensions (16:9 widescreen)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]
    
    for slide_data in SLIDES_DATA:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Add background styling
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(248, 250, 252) # Light steel grey background
        
        if slide_data.get("is_title"):
            # Title slide design
            # Top color band
            band = slide.shapes.add_shape(
                1, # rectangle
                0, 0, prs.slide_width, Inches(0.4)
            )
            band.fill.solid()
            band.fill.fore_color.rgb = NAVY
            band.line.color.rgb = NAVY
            
            # Text box
            txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(4))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p = tf.paragraphs[0]
            p.text = slide_data["title"]
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = NAVY
            p.font.name = "Arial"
            
            p2 = tf.add_paragraph()
            p2.text = slide_data["subtitle"]
            p2.font.size = Pt(20)
            p2.font.color.rgb = CHARCOAL
            p2.font.name = "Arial"
            p2.space_before = Pt(30)
            
        else:
            # Standard Content Slide
            # Slide Header
            headerBox = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.8), Inches(1))
            tf_h = headerBox.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = slide_data["title"]
            p_h.font.size = Pt(32)
            p_h.font.bold = True
            p_h.font.color.rgb = NAVY
            p_h.font.name = "Arial"
            
            # Bottom accent line under title
            line = slide.shapes.add_shape(
                1, # rectangle
                Inches(0.75), Inches(1.3), Inches(11.833), Inches(0.05)
            )
            line.fill.solid()
            line.fill.fore_color.rgb = STEEL_BLUE
            line.line.color.rgb = STEEL_BLUE
            
            # Content Box
            contentBox = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.8), Inches(4.8))
            tf_c = contentBox.text_frame
            tf_c.word_wrap = True
            
            for i, bullet in enumerate(slide_data["bullets"]):
                p_c = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
                p_c.text = "• " + bullet
                p_c.font.size = Pt(18)
                p_c.font.color.rgb = CHARCOAL
                p_c.font.name = "Arial"
                p_c.space_after = Pt(14)
                
    output_path = "/Users/manojvardhan/Bluestocks/reports/Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")
    
    # Copy to submission folder
    dest_path = "/Users/manojvardhan/Bluestocks/ManojVardhan_Submission/PPT / Slides/Presentation.pptx"
    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
    import shutil
    shutil.copy(output_path, dest_path)
    print(f"Presentation copied to submission folder: {dest_path}")

if __name__ == "__main__":
    create_deck()
